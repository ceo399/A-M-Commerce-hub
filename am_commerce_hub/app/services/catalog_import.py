"""メーカーカタログ取り込みパーサ（テンプレート形式 CSV/Excel 対応）。

「メーカーカタログ_取込テンプレート.xlsx」の列定義に合わせ、日本語ヘッダを
内部フィールドへ正規化し、catalog_service.sync_catalog が受け取れる行に変換する。
依存は標準ライブラリ(csv) と openpyxl のみ（pandas不要）。

返り値 ParseResult:
  rows     : sync_catalog にそのまま渡せる正規化済み行のリスト
  errors   : 取り込めなかった行（行番号＋理由）
  warnings : 補正・注意（行番号＋内容）
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

# 入力シート名（xlsx）。無ければ先頭シートを使う。
DATA_SHEET = "商品データ"

# 供給状況の語彙（表記ゆれを吸収）
ACTIVE_WORDS = {"供給中", "取扱中", "販売中", "在庫あり", "active"}
DISCONTINUED_WORDS = {"生産終了", "生産完了", "廃番", "廃盤", "終了", "discontinued"}

# ヘッダ → 内部フィールドの対応（spec.* はスペック辞書に格納）
HEADER_MAP = {
    "SKU": "sku",
    "商品名": "name",
    "JANコード": "jan",
    "メーカー名": "manufacturer",
    "ブランド名": "brand",
    "供給状況": "_supply",
    "在庫数": "_stock",
    "カテゴリ": "spec.category",
    "希望小売価格": "list_price",
    "仕入原価": "cost_price",
    "色": "spec.color",
    "サイズ/寸法": "spec.size",
    "重量": "spec.weight",
    "素材": "spec.material",
    "対応機種/用途": "spec.compatibility",
    "画像ファイル名/URL": "_images",
}


@dataclass
class ParseResult:
    rows: list[dict] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)
    warnings: list[dict] = field(default_factory=list)
    total_rows: int = 0

    def summary(self) -> dict:
        return {"total": self.total_rows, "ok": len(self.rows),
                "errors": len(self.errors), "warnings": len(self.warnings)}


# --------------------------------------------------------------------------
# 公開関数
# --------------------------------------------------------------------------
def parse_catalog(path: str | Path, sheet: str | None = None) -> ParseResult:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        header, records = _read_xlsx(path, sheet)
    elif suffix in (".csv", ".tsv", ".txt"):
        header, records = _read_csv(path, delimiter="\t" if suffix == ".tsv" else ",")
    else:
        raise ValueError(f"未対応のファイル形式です: {suffix}（CSV または xlsx）")
    return _normalize(header, records)


# --------------------------------------------------------------------------
# 読み込み（フォーマット別）
# --------------------------------------------------------------------------
def _read_csv(path: Path, delimiter: str = ",") -> tuple[list[str], list[list[str]]]:
    # utf-8-sig で BOM を吸収。失敗時は cp932(Shift_JIS) を試す。
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp932"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")
    reader = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    reader = [r for r in reader if any((c or "").strip() for c in r)]  # 空行除去
    if not reader:
        return [], []
    return [c.strip() for c in reader[0]], reader[1:]


def _read_xlsx(path: Path, sheet: str | None) -> tuple[list[str], list[list[str]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    name = sheet or (DATA_SHEET if DATA_SHEET in wb.sheetnames else wb.sheetnames[0])
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    rows = [r for r in rows if any(c not in (None, "") for c in r)]  # 空行除去
    if not rows:
        return [], []
    header = [(str(c).strip() if c is not None else "") for c in rows[0]]
    body = [[("" if c is None else c) for c in r] for r in rows[1:]]
    return header, body


# --------------------------------------------------------------------------
# 正規化・検証
# --------------------------------------------------------------------------
def _normalize(header: list[str], records: list[list[str]]) -> ParseResult:
    result = ParseResult(total_rows=len(records))
    if not header:
        result.errors.append({"row": 0, "message": "ヘッダ行が見つかりません"})
        return result

    # 特徴列・未知列を把握
    feature_idx = [i for i, h in enumerate(header) if h.startswith("特徴")]
    known = set(HEADER_MAP) | {header[i] for i in feature_idx}
    seen_sku: set[str] = set()

    for n, record in enumerate(records, start=2):  # 2行目=最初のデータ行
        def cell(col_name: str) -> str:
            try:
                i = header.index(col_name)
            except ValueError:
                return ""
            return _s(record[i]) if i < len(record) else ""

        sku = cell("SKU")
        if not sku:
            result.errors.append({"row": n, "message": "SKUが空です（必須）"})
            continue
        if sku in seen_sku:
            result.warnings.append({"row": n, "message": f"SKU重複のためスキップ: {sku}"})
            continue
        seen_sku.add(sku)

        name = cell("商品名")
        if not name:
            result.errors.append({"row": n, "message": f"商品名が空です（必須）: {sku}"})
            continue

        # 供給状況 → 廃盤判定
        supply = cell("供給状況")
        stock = _to_int(cell("在庫数"))
        if not supply:
            result.warnings.append({"row": n, "message": f"供給状況が空。供給中とみなします: {sku}"})
            discontinued = False
        elif supply in ACTIVE_WORDS:
            discontinued = False
        elif supply in DISCONTINUED_WORDS:
            discontinued = True
        else:
            result.warnings.append({"row": n, "message": f"供給状況『{supply}』を判定できず供給中とみなします: {sku}"})
            discontinued = False
        # 生産終了でも在庫が残っていれば販売継続（在庫0で初めて廃盤シグナル）
        in_stock = not (discontinued and (stock or 0) <= 0)

        # スペック構築
        spec: dict = {}
        for h, target in HEADER_MAP.items():
            if target.startswith("spec."):
                v = cell(h)
                if v:
                    spec[target.split(".", 1)[1]] = v
        features = [_s(record[i]) for i in feature_idx if i < len(record) and _s(record[i])]
        if features:
            spec["features"] = features
        images_raw = cell("画像ファイル名/URL")
        if images_raw:
            spec["images"] = [x.strip() for x in images_raw.replace("；", ";").split(";") if x.strip()]
        # 未知列もスペックとして保全
        for i, h in enumerate(header):
            if h and h not in known and i < len(record):
                v = _s(record[i])
                if v:
                    spec[h] = v

        row = {
            "sku": sku, "name": name,
            "jan": cell("JANコード") or None,
            "manufacturer": cell("メーカー名") or None,
            "brand": cell("ブランド名") or None,
            "in_stock": in_stock,
            "stock_qty": stock,
            "list_price": _to_float(cell("希望小売価格")),
            "cost_price": _to_float(cell("仕入原価")),
            "spec": spec,
        }
        result.rows.append(row)

    return result


# --------------------------------------------------------------------------
# 小物
# --------------------------------------------------------------------------
def _s(v) -> str:
    return "" if v is None else str(v).strip()


def _to_int(v: str) -> int | None:
    v = _digits(v)
    return int(v) if v else None


def _to_float(v: str) -> float | None:
    v = _digits(v, dot=True)
    return float(v) if v else None


def _digits(v: str, dot: bool = False) -> str:
    """¥・,・円・全角空白などを除去して数値文字だけ残す。"""
    if not v:
        return ""
    allowed = "0123456789" + ("." if dot else "")
    return "".join(ch for ch in str(v) if ch in allowed)
